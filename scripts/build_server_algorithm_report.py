"""Tạo báo cáo Word kiến trúc và điều phối hệ thống AGV từ mã và tài liệu hiện hành."""

from __future__ import annotations

import json
import math
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
FIG = ROOT / "docs" / "figures" / "server-report"
OUT = ROOT / "docs" / "exports" / "Bao_cao_thuat_toan_server_AGV.docx"
FIG.mkdir(parents=True, exist_ok=True)
OUT.parent.mkdir(parents=True, exist_ok=True)

BLACK = "#000000"
DARK = "#303030"
MID = "#666666"
GRID = "#B7B7B7"
LIGHT = "#F2F2F2"
PALE = "#FAFAFA"
WHITE = "#FFFFFF"
FONT = Path("C:/Windows/Fonts/times.ttf")
FONT_BOLD = Path("C:/Windows/Fonts/timesbd.ttf")
FONT_ITALIC = Path("C:/Windows/Fonts/timesi.ttf")


def f(size: int, bold: bool = False, italic: bool = False):
    path = FONT_BOLD if bold else FONT_ITALIC if italic else FONT
    return ImageFont.truetype(str(path), size)


def canvas(width: int, height: int):
    image = Image.new("RGB", (width, height), WHITE)
    return image, ImageDraw.Draw(image)


def wrapped_lines(draw, value: str, width: int, font):
    words = value.split()
    lines, current = [], ""
    for word in words:
        candidate = f"{current} {word}".strip()
        if current and draw.textbbox((0, 0), candidate, font=font)[2] > width:
            lines.append(current)
            current = word
        else:
            current = candidate
    if current:
        lines.append(current)
    return lines


def text_lines(draw, xy, value, width, size=28, color=BLACK, bold=False, align="center", gap=6):
    font = f(size, bold)
    lines = wrapped_lines(draw, value, width, font)
    x, y = xy
    for line in lines:
        bbox = draw.textbbox((0, 0), line, font=font)
        line_width = bbox[2] - bbox[0]
        xx = x - line_width / 2 if align == "center" else x
        draw.text((xx, y), line, font=font, fill=color)
        y += size + gap
    return y


def box(draw, rect, title, subtitle="", fill=PALE, border=GRID, title_size=31, sub_size=23):
    x0, y0, x1, y1 = rect
    draw.rounded_rectangle(rect, radius=16, fill=fill, outline=border, width=3)
    if subtitle:
        text_lines(draw, ((x0 + x1) / 2, y0 + 24), title, x1 - x0 - 30, title_size, BLACK, True)
        text_lines(draw, ((x0 + x1) / 2, y0 + 74), subtitle, x1 - x0 - 34, sub_size, DARK)
    else:
        lines = wrapped_lines(draw, title, x1 - x0 - 34, f(title_size, True))
        top = (y0 + y1) / 2 - len(lines) * (title_size + 5) / 2
        text_lines(draw, ((x0 + x1) / 2, top), title, x1 - x0 - 34, title_size, BLACK, True)


def arrow(draw, start, end, color=DARK, width=5, head=16):
    x0, y0 = start
    x1, y1 = end
    draw.line((x0, y0, x1, y1), fill=color, width=width)
    angle = math.atan2(y1 - y0, x1 - x0)
    p1 = (x1 - head * math.cos(angle - 0.48), y1 - head * math.sin(angle - 0.48))
    p2 = (x1 - head * math.cos(angle + 0.48), y1 - head * math.sin(angle + 0.48))
    draw.polygon([(x1, y1), p1, p2], fill=color)


def diagram_title(draw, value):
    draw.text((65, 35), value, font=f(38, True), fill=BLACK)


def architecture():
    image, draw = canvas(1800, 800)
    diagram_title(draw, "Kiến trúc tổng thể hệ thống AGV")
    blocks = [
        (70, 320, "Người vận hành", "Trình duyệt và dashboard"),
        (385, 700, "Server web", "Node.js gateway và WebSocket"),
        (770, 1090, "Server quản lý", "C++ Core điều phối tập trung"),
        (1160, 1430, "MQTT broker", "Kênh trao đổi với robot"),
        (1500, 1740, "AGV 01 và 02", "ESP32, RFID, cảm biến, motor"),
    ]
    for x0, x1, a, b in blocks:
        box(draw, (x0, 190, x1, 405), a, b, fill=LIGHT, border=MID)
    links = [(320, 385, "HTTP và WS"), (700, 770, "WS v1"), (1090, 1160, "MQTT"), (1430, 1500, "WiFi")]
    for x0, x1, label in links:
        arrow(draw, (x0 + 8, 270), (x1 - 8, 270), DARK, 5)
        arrow(draw, (x1 - 8, 325), (x0 + 8, 325), MID, 4)
        text_lines(draw, ((x0 + x1) / 2, 225), label, x1 - x0 + 70, 18, BLACK, True)
    box(draw, (410, 535, 675, 690), "SQLite", "Đơn hàng, lịch sử, kho và báo cáo", fill=WHITE, border=MID, title_size=29, sub_size=21)
    arrow(draw, (520, 405), (520, 525), DARK, 5)
    arrow(draw, (570, 525), (570, 405), MID, 4)
    box(draw, (795, 535, 1065, 690), "Trạng thái Core", "Bản đồ, task, route, mutex và lift", fill=WHITE, border=MID, title_size=29, sub_size=21)
    arrow(draw, (930, 405), (930, 525), DARK, 5)
    draw.text((70, 740), "Sơ đồ khối chức năng; các cổng mặc định là gateway 3000, Core 8080 và MQTT 1883.", font=f(23, italic=True), fill=DARK)
    image.save(FIG / "01-kien-truc.png")


def core_modules():
    image, draw = canvas(1800, 950)
    diagram_title(draw, "Các module chính của C++ Core")
    modules = [
        (90, 175, 500, 360, "WebCommands", "Nhận lệnh WS và trả command_result"),
        (695, 175, 1105, 360, "FleetManager", "Queue, phân công, route, task và lift"),
        (1300, 175, 1710, 360, "MQTT callback", "Telemetry, diagnostic, ACK và mutex"),
        (90, 525, 500, 710, "MapGraph", "Đồ thị có hướng và Dijkstra"),
        (695, 525, 1105, 710, "MutexArbiter", "Owner, hàng chờ và grant token"),
        (1300, 525, 1710, 710, "RuntimeStore", "Checkpoint, counter và phục hồi"),
    ]
    for x0, y0, x1, y1, title, subtitle in modules:
        box(draw, (x0, y0, x1, y1), title, subtitle, fill=LIGHT if y0 < 500 else WHITE, border=MID, title_size=30, sub_size=22)
    arrow(draw, (500, 265), (685, 265), DARK, 5)
    arrow(draw, (1300, 265), (1115, 265), DARK, 5)
    arrow(draw, (900, 370), (900, 515), DARK, 5)
    arrow(draw, (695, 620), (510, 620), MID, 4)
    arrow(draw, (1105, 620), (1290, 620), MID, 4)
    box(draw, (455, 805, 1345, 905), "Nguyên tắc phát lệnh", "Lưu checkpoint trước khi phát command hoặc grant", fill=PALE, border=BLACK, title_size=29, sub_size=22)
    image.save(FIG / "02-ba-luong-server.png")


def gateway_modules():
    image, draw = canvas(1800, 900)
    diagram_title(draw, "Server web Node.js và SQLite")
    box(draw, (90, 170, 480, 360), "HTTP API", "Xác thực, kiểm tra dữ liệu và truy vấn", fill=LIGHT, border=MID)
    box(draw, (705, 170, 1095, 360), "Gateway", "Ghép request_id với kết quả Core", fill=WHITE, border=BLACK)
    box(draw, (1320, 170, 1710, 360), "WebSocket", "State, state_delta và trạng thái kết nối", fill=LIGHT, border=MID)
    arrow(draw, (480, 265), (695, 265), DARK, 5)
    arrow(draw, (1095, 265), (1310, 265), DARK, 5)
    arrow(draw, (1310, 315), (1105, 315), MID, 4)
    services = [
        (95, 510, 470, 710, "Đơn hàng", "Đặt chỗ và vòng đời chuyển kệ"),
        (520, 510, 895, 710, "Kho", "K1 đến K6, kệ trên xe và đối soát"),
        (945, 510, 1320, 710, "Lịch sử", "Task, robot, log và KPI"),
        (1370, 510, 1710, 710, "Báo cáo ca", "Ảnh đầu cuối và sự kiện bất biến"),
    ]
    for x0, y0, x1, y1, title, subtitle in services:
        box(draw, (x0, y0, x1, y1), title, subtitle, fill=PALE, border=GRID, title_size=29, sub_size=21)
        arrow(draw, ((x0 + x1) / 2, 500), ((x0 + x1) / 2, 370), MID, 4)
    draw.rounded_rectangle((250, 785, 1550, 855), radius=14, fill=LIGHT, outline=MID, width=3)
    text_lines(draw, (900, 805), "SQLite transaction giữ thay đổi nghiệp vụ và bằng chứng trong cùng một lần ghi", 1200, 26, BLACK, True)
    image.save(FIG / "03-phan-cong-nhiem-vu.png")


def warehouse_map():
    data = json.loads((ROOT / "data" / "map_layout.json").read_text(encoding="utf-8"))
    image, draw = canvas(1800, 1120)
    diagram_title(draw, "Bản đồ kho")
    boundary = (85, 115, 1715, 920)
    draw.rounded_rectangle(boundary, radius=24, fill=WHITE, outline=BLACK, width=4)
    draw.rectangle((115, 102, 405, 135), fill=WHITE)
    draw.text((125, 105), "PHẠM VI SƠ ĐỒ KHO", font=f(22, True), fill=BLACK)
    xs = [float(node["x"]) for node in data["nodes"]]
    ys = [float(node["y"]) for node in data["nodes"]]
    min_x, max_x = min(xs), max(xs)
    min_y, max_y = min(ys), max(ys)

    def map_point(x, y):
        px = 175 + (float(x) - min_x) / (max_x - min_x) * 1450
        py = 205 + (float(y) - min_y) / (max_y - min_y) * 625
        return px, py

    points = {node["id"]: map_point(node["x"], node["y"]) for node in data["nodes"]}
    for edge in data["edges"]:
        a, b = points[edge["from"]], points[edge["to"]]
        via = (edge.get("geometry") or {}).get("via")
        if isinstance(via, dict) and "x" in via and "y" in via:
            control = map_point(via["x"], via["y"])
            curve = []
            for step in range(30):
                t = step / 29
                x = (1 - t) ** 2 * a[0] + 2 * (1 - t) * t * control[0] + t**2 * b[0]
                y = (1 - t) ** 2 * a[1] + 2 * (1 - t) * t * control[1] + t**2 * b[1]
                curve.append((x, y))
            draw.line(curve, fill=GRID, width=4)
            arrow(draw, curve[-3], curve[-1], MID, 3, 11)
            continue
        vx, vy = b[0] - a[0], b[1] - a[1]
        length = max(1, math.hypot(vx, vy))
        nx, ny = -vy / length * 5, vx / length * 5
        start = (a[0] + vx / length * 25 + nx, a[1] + vy / length * 25 + ny)
        end = (b[0] - vx / length * 29 + nx, b[1] - vy / length * 29 + ny)
        arrow(draw, start, end, GRID, 3, 11)
    for node in data["nodes"]:
        node_id = node["id"]
        x, y = points[node_id]
        node_type = node["type"]
        if node_type == "station":
            draw.rounded_rectangle((x - 27, y - 21, x + 27, y + 21), radius=7, fill=WHITE, outline=BLACK, width=5)
        elif node_type == "charging":
            draw.rectangle((x - 23, y - 23, x + 23, y + 23), fill=WHITE, outline=BLACK, width=5)
        elif node_type == "mutex":
            draw.ellipse((x - 25, y - 25, x + 25, y + 25), fill=WHITE, outline=BLACK, width=5)
            draw.ellipse((x - 19, y - 19, x + 19, y + 19), outline=MID, width=2)
        else:
            draw.ellipse((x - 21, y - 21, x + 21, y + 21), fill=WHITE, outline=MID, width=4)
        draw.text((x, y - 1), node_id, font=f(22, True), fill=BLACK, anchor="mm")
    legend = [(160, "K", "Vị trí kho K1 đến K6"), (575, "M", "Vùng giao nhau M1 đến M6"), (1050, "S", "Trạm S1 và S2"), (1395, "G", "Node nối tuyến")]
    for x, symbol, label in legend:
        draw.rounded_rectangle((x, 970, x + 45, 1015), radius=7, fill=WHITE, outline=BLACK, width=3)
        draw.text((x + 22, 992), symbol, font=f(22, True), fill=BLACK, anchor="mm")
        draw.text((x + 60, 975), label, font=f(23), fill=BLACK)
    draw.text((85, 1060), f"Nguồn: data/map_layout.json, revision {data.get('revision', '?')}, {len(data['nodes'])} node và {len(data['edges'])} cạnh. Tọa độ dùng để minh họa, chưa phải kích thước vật lý đã đo.", font=f(21, italic=True), fill=DARK)
    image.save(FIG / "04-ban-do-kho.png")


def dashboard_wireframe():
    image, draw = canvas(1800, 980)
    diagram_title(draw, "Bố cục dashboard giám sát")
    draw.rounded_rectangle((75, 125, 1725, 900), radius=20, fill=WHITE, outline=BLACK, width=4)
    draw.rectangle((75, 125, 1725, 220), fill=LIGHT, outline=BLACK, width=3)
    draw.text((120, 155), "Web", font=f(28, True), fill=BLACK)
    for x, label in [(1180, "Core"), (1370, "MQTT"), (1565, "Dữ liệu")]:
        draw.ellipse((x, 158, x + 22, 180), fill=WHITE, outline=BLACK, width=3)
        draw.text((x + 35, 152), label, font=f(22), fill=BLACK)
    for x, label in [(120, "Tổng quan"), (395, "Nhiệm vụ"), (670, "Bản đồ"), (945, "Báo cáo")]:
        draw.rounded_rectangle((x, 245, x + 230, 305), radius=12, fill=LIGHT if label == "Tổng quan" else WHITE, outline=MID, width=2)
        draw.text((x + 115, 275), label, font=f(24, True), fill=BLACK, anchor="mm")
    box(draw, (115, 345, 560, 565), "Đội xe", "RFID, pin, lỗi, tuổi dữ liệu và task", fill=PALE, border=GRID, title_size=30, sub_size=22)
    box(draw, (600, 345, 1250, 690), "Bản đồ và tuyến", "Node, cạnh, vùng mutex, vật cản và marker robot", fill=WHITE, border=MID, title_size=31, sub_size=23)
    box(draw, (1290, 345, 1685, 565), "Điều khiển", "Mở đóng ca, giao việc và thao tác kho", fill=PALE, border=GRID, title_size=30, sub_size=22)
    box(draw, (115, 610, 560, 835), "Trạng thái kho", "Kệ tại K1 đến K6, trên xe và đối soát", fill=PALE, border=GRID, title_size=30, sub_size=22)
    box(draw, (1290, 610, 1685, 835), "Cảnh báo", "Offline, stale, unknown và needs_attention", fill=PALE, border=GRID, title_size=30, sub_size=22)
    draw.text((95, 925), "Khung dây minh họa cách phân khu chức năng, không phải ảnh chụp giao diện.", font=f(22, italic=True), fill=DARK)
    image.save(FIG / "05-ack-va-rfid.png")


def sequence_flow():
    image, draw = canvas(1800, 1250)
    diagram_title(draw, "Sơ đồ tuần tự một lượt chuyển kệ")
    lanes = [(180, "Người vận hành"), (510, "Gateway"), (850, "C++ Core"), (1190, "Robot"), (1530, "SQLite")]
    for x, label in lanes:
        box(draw, (x - 125, 125, x + 125, 220), label, fill=LIGHT, border=MID, title_size=25)
        draw.line((x, 220, x, 1170), fill=GRID, width=3)
    steps = [
        (285, 180, 510, "Tạo đơn nguồn và đích"), (390, 510, 1530, "Đặt chỗ và lưu đơn"),
        (495, 510, 850, "add_task với request_id"), (600, 850, 1190, "GOTO có cmd_id qua MQTT"),
        (705, 1190, 850, "ACK rồi telemetry RFID đích"), (810, 850, 510, "task completed theo task_id"),
        (915, 510, 850, "warehouse_lift"), (1020, 850, 1190, "LIFT rồi DONE_UP hoặc DONE_DOWN"),
        (1125, 510, 1530, "Cập nhật sổ kho trong transaction"),
    ]
    for y, x0, x1, label in steps:
        arrow(draw, (x0, y), (x1 - 14 if x1 > x0 else x1 + 14, y), DARK, 4, 14)
        text_lines(draw, ((x0 + x1) / 2, y - 45), label, abs(x1 - x0) - 35, 20, BLACK, True)
    draw.text((70, 1205), "ACK chỉ xác nhận robot đã nhận lệnh; RFID và kết quả servo tương quan mới cho phép tiến sang bước tiếp theo.", font=f(22, italic=True), fill=DARK)
    image.save(FIG / "06-phuc-hoi.png")


def safety_flow():
    image, draw = canvas(1800, 900)
    diagram_title(draw, "Phản ứng khi dữ liệu hoặc liên kết không còn tin cậy")
    events = [
        (85, 175, 515, 355, "Mất MQTT hoặc telemetry", "Không kết luận robot đã dừng hoặc rời vùng"),
        (685, 175, 1115, 355, "Core hoặc gateway khởi động lại", "Khôi phục dữ liệu bền vững nhưng xác minh lại vị trí"),
        (1285, 175, 1715, 355, "Kết quả lệnh chưa xác định", "Không tự phát lại thao tác có thể đã chạy"),
    ]
    for x0, y0, x1, y1, title, subtitle in events:
        box(draw, (x0, y0, x1, y1), title, subtitle, fill=LIGHT, border=MID, title_size=27, sub_size=21)
        arrow(draw, ((x0 + x1) / 2, 365), ((x0 + x1) / 2, 485), DARK, 5)
    box(draw, (330, 500, 1470, 665), "Trạng thái an toàn thận trọng", "Giữ task, route hoặc quyền vùng chưa xác minh; khóa điều khiển và yêu cầu đối soát", fill=WHITE, border=BLACK, title_size=31, sub_size=23)
    for x, label in [(570, "Telemetry RFID mới"), (900, "Đối soát vật lý"), (1230, "Lệnh mới có ID mới")]:
        arrow(draw, (x, 675), (x, 750), MID, 4)
        text_lines(draw, (x, 770), label, 280, 23, BLACK, True)
    image.save(FIG / "07-luong-don-kho.png")


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    element = OxmlElement("w:tblHeader")
    element.set(qn("w:val"), "true")
    tr_pr.append(element)


def set_borders(table):
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement("w:" + edge)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "5")
        element.set(qn("w:color"), "B7B7B7")
        borders.append(element)
    table._tbl.tblPr.append(borders)


def shade(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    element = OxmlElement("w:shd")
    element.set(qn("w:fill"), color)
    tc_pr.append(element)


def cell_margin(cell, top=100, start=115, bottom=100, end=115):
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        tag = "w:" + side
        node = margins.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_run_font(run, size=None, bold=None, italic=None):
    run.font.name = "Times New Roman"
    run_properties = run._element.get_or_add_rPr()
    fonts = run_properties.get_or_add_rFonts()
    for key in ("w:ascii", "w:hAnsi", "w:eastAsia"):
        fonts.set(qn(key), "Times New Roman")
    run.font.color.rgb = RGBColor(0, 0, 0)
    color = run_properties.get_or_add_color()
    color.set(qn("w:val"), "000000")
    for attribute in ("w:themeColor", "w:themeTint", "w:themeShade"):
        qualified = qn(attribute)
        if qualified in color.attrib:
            del color.attrib[qualified]
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    if widths:
        table_width = table._tbl.tblPr.find(qn("w:tblW"))
        if table_width is None:
            table_width = OxmlElement("w:tblW")
            table._tbl.tblPr.append(table_width)
        table_width.set(qn("w:w"), str(int(sum(widths) * 567)))
        table_width.set(qn("w:type"), "dxa")
        for index, width in enumerate(widths):
            table.columns[index].width = Cm(width)
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        if widths:
            cell.width = Cm(widths[index])
        cell.text = header
        shade(cell, "D9D9D9")
        cell_margin(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for paragraph in cell.paragraphs:
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                set_run_font(run, 10.5, bold=True)
    set_repeat_header(table.rows[0])
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = str(value)
            if widths:
                cells[index].width = Cm(widths[index])
            if row_index % 2:
                shade(cells[index], "F7F7F7")
            cell_margin(cells[index])
            cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cells[index].paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.1
                for run in paragraph.runs:
                    set_run_font(run, 10.5)
    set_borders(table)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_para(doc, text, indent=True):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.line_spacing = 1.3
    paragraph.paragraph_format.space_after = Pt(5)
    if indent:
        paragraph.paragraph_format.first_line_indent = Cm(1.0)
    run = paragraph.add_run(text)
    set_run_font(run, 13)
    return paragraph


def add_pseudocode(doc, lines):
    for line in lines:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Cm(1.2)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0
        run = paragraph.add_run(line)
        run.font.name = "Courier New"
        run.font.size = Pt(10.5)
        run.font.color.rgb = RGBColor(0, 0, 0)
        run_properties = run._element.get_or_add_rPr()
        fonts = run_properties.get_or_add_rFonts()
        for key in ("w:ascii", "w:hAnsi", "w:eastAsia"):
            fonts.set(qn(key), "Courier New")
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)


def add_figure(doc, filename, caption, width=16.0):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(7)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.keep_with_next = True
    paragraph.add_run().add_picture(str(FIG / filename), width=Cm(width))
    caption_paragraph = doc.add_paragraph(style="Caption")
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_run = caption_paragraph.add_run(caption)
    set_run_font(caption_run, 11, italic=True)


def add_page_num(paragraph):
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)


def configure_style(style, size, bold=False, italic=False):
    style.font.name = "Times New Roman"
    run_properties = style._element.get_or_add_rPr()
    fonts = run_properties.get_or_add_rFonts()
    for key in ("w:ascii", "w:hAnsi", "w:eastAsia"):
        fonts.set(qn(key), "Times New Roman")
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = italic
    style.font.color.rgb = RGBColor(0, 0, 0)
    color = run_properties.get_or_add_color()
    color.set(qn("w:val"), "000000")
    for attribute in ("w:themeColor", "w:themeTint", "w:themeShade"):
        qualified = qn(attribute)
        if qualified in color.attrib:
            del color.attrib[qualified]


def build_doc():
    doc = Document()
    section = doc.sections[0]
    section.page_height = Cm(29.7)
    section.page_width = Cm(21)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(2.0)
    section.header_distance = Cm(1.0)
    section.footer_distance = Cm(1.0)
    styles = doc.styles
    configure_style(styles["Normal"], 13)
    styles["Normal"].paragraph_format.line_spacing = 1.3
    styles["Normal"].paragraph_format.space_after = Pt(5)
    configure_style(styles["Title"], 19, bold=True)
    styles["Title"].paragraph_format.space_after = Pt(10)
    styles["Title"].paragraph_format.keep_with_next = True
    title_ppr = styles["Title"]._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)
    configure_style(styles["Subtitle"], 13, italic=True)
    styles["Subtitle"].paragraph_format.space_after = Pt(8)
    configure_style(styles["Heading 1"], 15, bold=True)
    styles["Heading 1"].paragraph_format.space_before = Pt(15)
    styles["Heading 1"].paragraph_format.space_after = Pt(7)
    styles["Heading 1"].paragraph_format.keep_with_next = True
    configure_style(styles["Heading 2"], 13, bold=True)
    styles["Heading 2"].paragraph_format.space_before = Pt(10)
    styles["Heading 2"].paragraph_format.space_after = Pt(4)
    styles["Heading 2"].paragraph_format.keep_with_next = True
    configure_style(styles["Caption"], 11, italic=True)
    styles["Caption"].paragraph_format.space_after = Pt(9)
    section.footer.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_num(section.footer.paragraphs[0])

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(title.add_run("BÁO CÁO KIẾN TRÚC VÀ THUẬT TOÁN ĐIỀU PHỐI AGV TRONG KHO"), 19, bold=True)
    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(subtitle.add_run("Phân tích theo mã nguồn và tài liệu hệ thống hiện hành"), 13, italic=True)
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(meta.add_run("Ngày đối chiếu 20 tháng 10 năm 2026"), 13)

    doc.add_heading("Giới thiệu bài toán", level=1)
    add_para(doc, "Kho sử dụng nhiều robot tự hành cần một điểm điều phối tập trung để phân công nhiệm vụ, chọn tuyến và kiểm soát quyền đi qua khu vực giao nhau. Nếu từng xe tự quyết định độc lập, hai xe có thể cùng tiến vào ngã tư, giữ tuyến chồng lấn hoặc tiếp tục chạy khi trạng thái của xe còn lại không còn được xác minh.")
    add_para(doc, "Dự án hiện quản lý hai AGV nhưng kiến trúc tách thành Core điều phối, gateway web, cơ sở dữ liệu, dashboard và kênh MQTT. Core sở hữu quyết định chuyển động và mutex; gateway sở hữu nghiệp vụ đơn hàng, sổ kho và báo cáo; giao diện giúp người vận hành theo dõi bản đồ, nhiệm vụ, kết nối và lỗi. Báo cáo phân biệt nội dung đã đối chiếu từ mã, hướng phát triển đề xuất và hạng mục chưa có số đo hoặc chưa nghiệm thu; không xem đề xuất là chức năng đã triển khai.")

    doc.add_heading("1 Kiến trúc tổng thể", level=1)
    add_para(doc, "Hệ thống dùng mô hình điều phối tập trung. Trình duyệt trao đổi với gateway Node.js qua HTTP và WebSocket. Gateway chuyển lệnh có request_id sang C++ Core qua WebSocket v1. Core quản lý bản đồ, đội xe, nhiệm vụ và mutex, sau đó gửi command JSON tới robot qua MQTT. Robot trả ACK, telemetry RFID, pin, vật cản và diagnostic. SQLite nằm phía gateway để lưu dữ liệu nghiệp vụ và lịch sử.")
    add_figure(doc, "01-kien-truc.png", "Hình 1. Sơ đồ khối kiến trúc tổng thể của hệ thống AGV.")
    add_table(doc, ["Thành phần", "Trách nhiệm chính", "Dữ liệu quyết định"], [
        ("C++ Core", "Tìm đường, phân công, theo dõi task, mutex, lift và phục hồi", "Map đang chạy, task_id, route, owner và grant_id"),
        ("Node.js gateway", "API web, tương quan kết quả Core và điều phối nghiệp vụ kho", "request_id, đơn hàng, trạng thái đối soát"),
        ("SQLite", "Lưu history, KPI, đơn hàng, sổ kệ và báo cáo ca", "Bản ghi nghiệp vụ đã xác nhận"),
        ("Dashboard", "Giám sát và gửi yêu cầu của người vận hành", "State nhận từ gateway cùng tuổi dữ liệu"),
        ("Robot", "Đọc cảm biến, điều khiển motor và thực hiện command", "RFID, cảm biến, PWM, encoder và trạng thái servo"),
    ], [3.2, 7.0, 6.8])

    doc.add_heading("2 Server quản lý (C++ Core)", level=1)
    add_para(doc, "Core được xây bằng C++17 và libmosquitto. Một tiến trình thống nhất có thể chạy với phạm vi tham gia (scope) chỉ xe 01, chỉ xe 02 hoặc cả hai; scope là tập ID robot mà tiến trình được phép nhận telemetry và giao nhiệm vụ. Các luồng lệnh web, sự kiện MQTT và nhịp điều phối cùng tác động đến trạng thái dùng chung; khóa nội bộ bảo vệ dữ liệu trong tiến trình, còn mutex zone bảo vệ khu vực giao nhau ngoài thực địa.")
    add_figure(doc, "02-ba-luong-server.png", "Hình 2. Quan hệ giữa các module chính của C++ Core.", width=14.5)
    doc.add_page_break()
    doc.add_heading("2.1 Tìm đường và phân công", level=2)
    add_para(doc, "MapGraph biểu diễn bản đồ kho bằng đồ thị có hướng G = (V, E). Bản hiện hành có 22 node và 35 cạnh. Mỗi cạnh e mang trọng số không âm w(e), được dùng như độ dài quy ước của đoạn đường; chi phí tuyến P là C(P) = tổng w(e) trên các cạnh thuộc P. Node có vật cản, node bị chặn bổ sung và cạnh bị khóa bị loại khỏi phép tìm đường. Node hoặc tuyến đang được xe khác giữ cũng được đưa vào tập cấm theo từng lượt xét, thay vì cộng một khoản phạt tắc nghẽn vào trọng số.")
    add_para(doc, "Dijkstra trả về tuyến khả thi có tổng trọng số nhỏ nhất trong tập cấm tại thời điểm tính. Một lần tìm đường có độ phức tạp O((|V| + |E|) log |V|) với hàng đợi ưu tiên và dùng O(|V| + |E|) bộ nhớ. Tọa độ x, y hiện phục vụ sơ đồ và chưa được hiệu chuẩn thành khoảng cách vật lý bảo đảm, nên chưa có heuristic cùng đơn vị và luôn không vượt chi phí thật để dùng A* một cách chặt chẽ.")
    add_para(doc, "FleetManager xét task theo priority giảm dần; cùng priority thì task có số thứ tự tạo nhỏ hơn được xét trước, tức FIFO trong cùng mức ưu tiên. Với mỗi task, Core lọc các xe đang thuộc scope, rảnh, còn kết nối, có telemetry mới, vị trí RFID hợp lệ và không ở trạng thái lỗi hoặc nâng hạ. Sau đó Core chạy tối đa hai lượt: lượt đầu tránh toàn bộ node bận và tuyến đang dùng; lượt sau dùng tập chặn cứng để tìm phương án an toàn còn lại. Xe có chi phí tuyến nhỏ nhất được chọn; nếu bằng nhau thì ID xe nhỏ hơn được chọn để kết quả xác định. Đây là phân công tham lam theo từng task, không phải tối ưu đồng thời toàn bộ hàng đợi.")
    add_para(doc, "Giả mã rút gọn của chu kỳ phân công:", indent=False)
    add_pseudocode(doc, [
        "for task theo (priority giảm dần, seq tăng dần):",
        "    best <- rỗng",
        "    for pass in [tránh node bận, tập chặn cứng]:",
        "        for agv in eligible(task):",
        "            p <- Dijkstra(agv.node, task.target, blocked)",
        "            nếu p tốt hơn best: best <- (agv, p, cost(p))",
        "        nếu best khác rỗng: break",
        "    nếu có best: gán task và tuyến",
        "    ngược lại: giữ task trong queue",
    ])
    add_para(doc, "Với T task đang chờ và R xe hợp lệ, cận trên của một chu kỳ phân công là O(T.R.(|V| + |E|) log |V|); hệ số hai đến từ hai lượt tập chặn được bỏ trong ký hiệu O lớn. Với bản đồ 22 node, 35 cạnh và tối đa hai xe, chi phí tính toán hiện nhỏ; rủi ro chính nằm ở chất lượng mô hình, bằng chứng chiếm dụng và chính sách công bằng.")
    add_table(doc, ["Phương pháp", "Khả năng", "Lý do chưa dùng hoặc vai trò hiện tại"], [
        ("Dijkstra hiện tại", "Tối ưu tổng trọng số tĩnh trên đồ thị có cạnh không âm", "Dễ kiểm chứng, xác định và phù hợp bản đồ nhỏ; không tối ưu thời gian chờ hoặc toàn đội xe."),
        ("A*", "Có thể giảm số node mở rộng khi có heuristic admissible", "Tọa độ hiện là sơ đồ, chưa đủ để suy ra cận dưới theo cùng đơn vị trọng số; cần hiệu chuẩn bản đồ trước."),
        ("Reservation table", "Đặt chỗ node và cạnh theo thời gian, xử lý tốt xung đột đối đầu", "Cần mô hình thời gian di chuyển, đồng hồ và cam kết lịch đáng tin cậy; các dữ liệu này chưa được nghiệm thu."),
        ("CBS", "Giải xung đột nhiều robot trong không gian-thời gian và có thể tối ưu theo giả thiết", "Độ phức tạp và yêu cầu mô hình cao hơn nhu cầu hai xe hiện tại; thời lượng chuyển động qua MQTT còn bất định."),
    ], [3.0, 5.0, 9.0])
    add_figure(doc, "04-ban-do-kho.png", "Hình 3. Bản đồ kho gồm các vị trí K, trạm S, node nối tuyến và vùng giao nhau M.")
    doc.add_heading("2.2 Tránh xung đột tại giao lộ", level=2)
    add_para(doc, "MutexArbiter chỉ cấp một owner cho mỗi vùng M1 đến M6, giữ hàng chờ FIFO riêng cho từng vùng và phát grant_id. ACK grant chỉ xác nhận robot đã nhận quyền. Quyền chỉ được trả khi robot cung cấp đúng token và bằng chứng rời vùng; mất telemetry hoặc MQTT không đủ để kết luận vùng đã trống. Firmware hiện dùng điều kiện encoder và cửa sổ RFID cho bằng chứng này, nhưng ngưỡng vẫn cần đo trên xe thật.")
    add_para(doc, "Để giảm nguy cơ hai xe giữ-chờ nhau trong hành lang gồm nhiều vùng mutex nối trực tiếp, Core xem toàn bộ thành phần mutex liên thông như một hành lang cần dành trước khi lập tuyến thay thế. Các node thường trên activePath của xe khác cũng được đưa vào tập chặn. Cơ chế này giảm hold-and-wait trong cấu hình hiện tại nhưng chưa tạo thành chứng minh deadlock-free cho mọi bản đồ hoặc mọi chuỗi cấp nhiều vùng rời nhau.")
    add_table(doc, ["Nguy cơ", "Cơ chế hiện có", "Khoảng trống còn lại"], [
        ("Deadlock", "Dành trước hành lang mutex liên thông, chặn tuyến đang chạy, giữ owner khi mất bằng chứng", "Chưa có thứ tự toàn cục cho mọi vùng rời nhau hoặc bộ phát hiện chu trình wait-for."),
        ("Livelock", "Tie-break theo ID, tuyến xác định và xe BLOCKED chỉ thử lại theo tick", "Chưa có bộ phát hiện dao động, backoff hoặc ngưỡng số lần replan."),
        ("Starvation", "FIFO trong từng priority và FIFO hàng chờ từng mutex", "Không có aging; task priority thấp có thể chờ lâu nếu liên tục xuất hiện task ưu tiên cao."),
        ("Đối đầu trên cạnh một làn", "Chặn node, activePath, uncertainPath và một phần tuyến hai chiều", "Chưa đặt chỗ cạnh theo thời gian; chưa chứng minh loại trừ mọi edge-swap hoặc xung đột hình học khi quay."),
    ], [3.0, 6.5, 7.5])
    add_para(doc, "Trọng số đường là tĩnh và không chứa độ dài hàng chờ mutex. Khi có vật cản, đường bị chặn hoặc xe khác mất kết nối trên tuyến, Core có thể tính lại tuyến; nếu chưa có tuyến thay thế, xe chuyển BLOCKED, nhận WAIT một lần và thử lại ở các tick sau. Vì vậy hệ thống có re-route theo sự kiện, nhưng chưa phải định tuyến nhận biết tắc nghẽn hay tối ưu makespan.")
    doc.add_heading("2.3 Lưu trạng thái và phục hồi", level=2)
    add_para(doc, "RuntimeStore lưu queue, fleet và mutex trước khi Core phát command hoặc grant. command_sequence.json và task_sequence.json phải tăng đơn điệu, không xóa hoặc dùng chung cho hai Core. Sau restart, task đang chạy chuyển sang interrupted, quyền vùng được giữ ở trạng thái cần phục hồi và vị trí cũ chưa được xem là vị trí mới cho tới khi có telemetry RFID hợp lệ.")
    add_para(doc, "Chính sách pin hiện có cấu hình bypass tạm để thử nghiệm và nhóm sạc thủ công. Pin vẫn được đo và hiển thị nhưng chưa chặn task theo ngưỡng trong cấu hình này. Đây là trạng thái thử nghiệm, không phải chính sách nghiệm thu.")

    doc.add_heading("3 Server web (Node.js gateway và SQLite)", level=1)
    add_para(doc, "Gateway dùng Node.js CommonJS, Express và thư viện ws. HTTP xử lý truy vấn hoặc thao tác có phản hồi; WebSocket chuyển lệnh sang Core và đẩy trạng thái về trình duyệt. Gateway chờ command_result có cùng request_id và core_session_id. Nếu liên kết đứt trong lúc chờ, kết quả được đánh dấu unknown hoặc needs_attention thay vì tự gửi lại thao tác có thể đã được thực hiện.")
    add_figure(doc, "03-phan-cong-nhiem-vu.png", "Hình 4. Các module nghiệp vụ của server web và lớp lưu trữ SQLite.")
    add_para(doc, "SQLite lưu lịch sử robot, task, KPI, snapshot, đơn hàng, trạng thái K1 đến K6, kệ trên xe và báo cáo ca. Quy trình chuyển kệ chỉ cập nhật sổ kho sau DONE_UP hoặc DONE_DOWN đúng command_id, đúng robot, đúng node và cargo tương ứng. Transaction giữ cập nhật vị trí kệ và log bằng chứng trong cùng một lần ghi.")
    add_table(doc, ["Tình huống", "Phản hồi của gateway"], [
        ("Core từ chối lệnh", "Trả lỗi theo command_result; không tạo thành công giả."),
        ("Hết thời gian hoặc mất Core khi chờ", "Trả kết quả chưa xác định; khóa retry tự động và yêu cầu đối soát."),
        ("Đổi core_session_id", "Đánh dấu khoảng quan sát và xử lý đơn đang chạy theo hướng an toàn."),
        ("SQLite ghi thất bại", "Rollback transaction; không công bố sổ kho đã đổi."),
        ("Dữ liệu robot cũ", "Giữ snapshot cuối để tham khảo nhưng gắn stale và khóa thao tác phụ thuộc dữ liệu mới."),
    ], [5.0, 12.0])

    doc.add_heading("4 Giao diện web (dashboard)", level=1)
    add_para(doc, "Dashboard hiện chia bốn khu vực Tổng quan, Nhiệm vụ, Bản đồ và Báo cáo. Tổng quan ưu tiên đội xe, trạng thái kết nối, bản đồ, tuyến và thao tác kho. Nhiệm vụ hiển thị task_id và lịch sử xe. Bản đồ quản lý vật cản, snapshot và editor. Báo cáo chứa KPI, nhật ký chuyển kệ và báo cáo đầu cuối ca.")
    add_figure(doc, "05-ack-va-rfid.png", "Hình 5. Khung chức năng chính của dashboard giám sát.")
    add_para(doc, "Giao diện hiển thị riêng trạng thái Web, Core và MQTT. Trạng thái agv_connected chỉ phản ánh liên kết gateway với Core; tuổi telemetry mới quyết định dữ liệu từng xe còn tin cậy hay không. Marker nội suy giữa hai RFID là ước lượng hiển thị, trong khi RFID gần nhất mới là mốc rời rạc dùng để xác nhận node và hoàn thành task.")
    add_para(doc, "Bảng trạng thái kho cho biết kệ tại K1 đến K6, kệ trên xe, vị trí đang được giữ chỗ và mục cần đối soát. Khi dữ liệu cũ, MQTT mất hoặc robot lỗi, dashboard dừng chỉ báo tiến trình đang chạy và khóa form giao việc. Giao diện không diễn giải HTTP 202 hoặc ACK thành robot đã hoàn thành.")
    add_table(doc, ["Khu vực", "Nội dung ưu tiên"], [
        ("Tổng quan", "Đội xe, trạng thái kết nối, tuyến đang chạy, kho và thao tác chính."),
        ("Nhiệm vụ", "Task hiện tại, task_id, lịch sử xe và kết quả hoàn thành hoặc gián đoạn."),
        ("Bản đồ", "Node, cạnh, vật cản, snapshot và công cụ chỉnh sửa có kiểm tra."),
        ("Báo cáo", "KPI, nhật ký, trạng thái kệ và báo cáo đầu cuối ca."),
    ], [4.0, 13.0])

    doc.add_heading("5 Luồng hoạt động tiêu biểu", level=1)
    add_para(doc, "Một lượt chuyển kệ bắt đầu khi người vận hành chọn vị trí nguồn có kệ và vị trí đích trống. Gateway đặt chỗ, lưu đơn rồi yêu cầu Core tạo task. Core chọn xe và tuyến, phát GOTO có cmd_id. Robot ACK lệnh và di chuyển; telemetry RFID tại đích cùng trạng thái phù hợp mới hoàn thành chặng.")
    add_figure(doc, "06-phuc-hoi.png", "Hình 6. Sơ đồ tuần tự một lượt chuyển kệ từ yêu cầu đến cập nhật SQLite.", width=15.6)
    add_para(doc, "Sau khi xe đến kho lấy, gateway gửi warehouse_lift. Chặng giao chỉ bắt đầu khi kết quả DONE_UP đúng lệnh xác nhận chu kỳ nâng đã kết thúc và cargo ảo ở trạng thái có hàng. Tại kho giao, DONE_DOWN tương quan mới cho phép cập nhật kệ sang vị trí đích. Cargo hiện là trạng thái phần mềm của chu kỳ servo, chưa phải bằng chứng có cảm biến tải hoặc ID kệ vật lý.")

    doc.add_heading("6 An toàn, bảo mật và xử lý lỗi", level=1)
    add_para(doc, "Thiết kế dùng nguyên tắc thất bại an toàn: khi trạng thái chuyển động hoặc quyền vùng không còn được xác minh, hệ thống giữ task, route hoặc owner thay vì giả định tài nguyên đã trống. MQTT đang dùng QoS 0, vì vậy ACK ứng dụng, retry cùng cmd_id, chống trùng và checkpoint giúp xử lý mất hoặc lặp bản tin nhưng không thay thế bằng chứng chuyển động vật lý.")
    add_figure(doc, "07-luong-don-kho.png", "Hình 7. Luồng xử lý thận trọng khi kết nối hoặc dữ liệu không còn tin cậy.")
    add_table(doc, ["Rủi ro", "Cơ chế đang có", "Giới hạn cần lưu ý"], [
        ("Hai xe cùng vào giao lộ", "Owner, hàng chờ, grant_id và release đúng token", "Ngưỡng rời vùng cần đo với thân xe và trượt bánh thật"),
        ("Mất MQTT hoặc telemetry", "Dừng tin cậy dữ liệu mới, giữ vùng và yêu cầu đối soát", "Không suy ra robot đã dừng chỉ từ mất tín hiệu"),
        ("Lặp hoặc lệch command ID", "Counter bền vững, dedup và STALE_COMMAND", "Clear error chỉ sau xác minh và dùng ID mới"),
        ("Lệnh web chưa rõ kết quả", "request_id, command_result và core_session_id", "Không replay tự động thao tác có tác động vật lý"),
        ("Truy cập trái phép", "Bearer token, kiểm tra Origin, HTTPS WSS và MQTT TLS ACL khi cấu hình", "Cần xác nhận cấu hình triển khai; mặc định phát triển chưa đồng nghĩa đã bật bảo mật"),
    ], [4.1, 6.7, 6.2])
    add_para(doc, "Cấu hình hardening hỗ trợ HTTPS hoặc WSS tại gateway và MQTT TLS cổng 8883 với tài khoản, ACL riêng cho Core và từng robot. Nếu Core và gateway chạy khác máy, liên kết Core cần mạng riêng, VPN hoặc TLS proxy vì socket Core hiện không tự cung cấp WSS. Chế độ không token chỉ phù hợp loopback.")
    add_para(doc, "Mô hình tập trung khiến C++ Core là điểm lỗi đơn đối với quyết định tuyến và quyền vùng. Khi Core ngừng hoạt động, thiết kế hiện ưu tiên giữ task và quyền chưa xác minh thay vì cho robot tự tiếp tục; cách này an toàn hơn nhưng làm giảm tính sẵn sàng. Hướng giảm thiểu gồm giám sát tiến trình, khởi động lại có RuntimeStore bền vững, broker và gateway được health-check, cùng quy trình đối soát trước khi cấp lại quyền. MQTT QoS 0 vẫn có thể làm mất bản tin, nên ACK ứng dụng và retry cùng cmd_id chỉ giảm rủi ro chứ không biến kênh thành exactly-once.")

    doc.add_heading("7 Kiểm thử và kết quả", level=1)
    add_para(doc, "Bằng chứng hiện có xác nhận logic phần mềm, hợp đồng giao tiếp và khả năng build. Các bài kiểm thử dùng broker, DB, map, counter và cổng riêng để không tác động dữ liệu vận hành. Kết quả dưới đây được đọc từ hồ sơ kiểm thử trong workspace; chúng không thay thế nghiệm thu hai robot thật.")
    add_table(doc, ["Lớp kiểm thử", "Kết quả đã ghi nhận", "Phạm vi bằng chứng"], [
        ("C++ Core", "Build và unit PASS; ca STALE_COMMAND, scope thống nhất, mutex và phục hồi có hồi quy", "Logic Core và file trạng thái trong môi trường test"),
        ("Gateway và SQLite", "npm test PASS 47 trên 47", "Gateway, security, orders, warehouse action và inventory với DB riêng"),
        ("Trình duyệt", "Edge headless PASS ở desktop 1440 và mobile 390", "Bốn khu vực, kho, stale, reconnect, báo cáo và thao tác form"),
        ("Bản đồ và route", "PASS 22 node; bốn cạnh cong ngoài; 60 tổ hợp ba node liên tiếp cho mỗi cấu hình firmware xe", "Tính nhất quán map giữa Core và firmware"),
        ("Firmware native", "Bộ hồi quy gần nhất PASS 107 trên 107", "RFID, line, mutex, safety và profile; chưa chạy xe thật"),
        ("Mô phỏng", "Các luồng hai xe, lỗi mạng và kho chạy trên S1: môi trường sự kiện rời rạc với robot ảo và MQTT riêng", "Hợp đồng và logic sự kiện; không mô phỏng động học hoặc vật lý"),
    ], [4.0, 5.2, 7.8])
    add_para(doc, "Trong bảng trên, bốn cạnh cong ngoài là bốn liên kết bo góc được vẽ ở biên bản đồ. Một tổ hợp ba node liên tiếp biểu diễn ngữ cảnh node trước - node hiện tại - node tiếp theo mà profile firmware dùng để kiểm tra quyết định rẽ. S1 là môi trường cô lập để kích sự kiện và kiểm tra hợp đồng, không phải tên một trạm sạc và không tạo bằng chứng về quãng đường phanh, trượt bánh hoặc tải cơ khí.")
    add_para(doc, "Kết luận kiểm thử hiện tại là hệ thống đã có bằng chứng phần mềm cho các đường đi chính và nhiều tình huống lỗi. Chưa có số đo nghiệm thu về khoảng phanh, độ trễ RFID, sai số encoder, thời gian rời mutex, độ ổn định bám line, tải servo hoặc xung đột hình học của hai xe trên sàn.")
    doc.add_heading("7.1 Kế hoạch đánh giá hiệu quả thuật toán", level=2)
    add_para(doc, "Các bộ PASS hiện chứng minh tính đúng hồi quy, chưa chứng minh chất lượng điều phối. Bộ đo dưới đây được xác định trước để tránh chọn chỉ số sau khi xem kết quả; hiện chưa có số liệu thực nghiệm đủ điều kiện đưa vào kết luận.")
    add_table(doc, ["Chỉ số", "Định nghĩa và cách thu"], [
        ("Makespan", "Từ lúc dispatch task đầu tiên đến khi task cuối cùng đạt trạng thái kết thúc trong cùng kịch bản."),
        ("Chờ mutex trung bình và P95", "Khoảng thời gian từ yêu cầu vùng đến grant; báo cáo trung bình, phân vị 95 và số mẫu."),
        ("Xung đột bị chặn", "Đếm lần yêu cầu bị xếp hàng, tuyến bị chặn hoặc xe chuyển BLOCKED do chiếm dụng."),
        ("Replan trên mỗi task", "Số lần Core tính lại tuyến sau vật cản, thay đổi chiếm dụng hoặc mất kết nối."),
        ("Thông lượng", "Số task hoàn thành trên một đơn vị thời gian với cùng bản đồ và tải nhiệm vụ."),
    ], [4.5, 12.5])
    add_para(doc, "Thí nghiệm đề xuất trên S1 gồm: (i) FIFO thuần so với priority hiện hành trên cùng chuỗi task; (ii) mutex hiện hành so với baseline không mutex chỉ trong mô phỏng cô lập để minh họa số xung đột, tuyệt đối không chạy baseline này trên robot thật; và (iii) tải cân bằng so với tải lệch về một khu vực. Mỗi kịch bản cần cùng seed, cùng danh sách task, ít nhất 30 lần lặp và báo cáo trung bình kèm P95. Chỉ sau khi thu log mới điền kết quả và kết luận hiệu quả.")

    doc.add_heading("8 Hạn chế và hướng phát triển", level=1)
    add_table(doc, ["Nội dung và bằng chứng hoàn thành"], [
        ("Bản đồ: Đo chiều dài cạnh, bán kính quay, độ rộng lối đi và vùng quét thân xe; cập nhật trọng số map rồi chạy lại route và xung đột.",),
        ("Pin: Đo điện áp và dung lượng theo tải, xác lập ngưỡng cảnh báo, dừng và sạc; bỏ bypass sau kiểm thử sạc an toàn.",),
        ("Servo: Thử không tải và có tải, đo góc, chu kỳ, dòng và lỗi; bổ sung cảm biến xác nhận kệ nếu cần.",),
        ("Nghiệm thu: Thực hiện theo thứ tự kiểm thử bàn, từng xe, hai xe tại giao lộ và chạy có tải; lưu biên bản cùng log truy vết.",),
        ("Điều phối nâng cao: Đo KPI mục 7.1; bổ sung aging, edge reservation hoặc định tuyến theo tắc nghẽn nếu số liệu cho thấy cần thiết.",),
        ("Độ sẵn sàng và bảo mật: Health-check Core, broker và gateway; diễn tập restart; bật TLS, ACL, quản lý bí mật và kiểm thử quyền truy cập.",),
    ], [17.0])
    add_para(doc, "Bản đồ, pin, servo và nghiệm thu là các nội dung còn thiếu để chuyển từ bằng chứng phần mềm sang bằng chứng vận hành thực tế. Trước khi hoàn thành, bản đồ chỉ là mô hình logic, chính sách pin còn bypass trong cấu hình thử, cargo chỉ là trạng thái phần mềm của chu kỳ servo và kết quả chưa phải nghiệm thu hệ thống trong kho. Điều phối nâng cao, độ sẵn sàng và bảo mật là hướng mở rộng sau khi có số liệu nền, không phải cam kết chức năng đã tồn tại.")

    doc.add_heading("Tài liệu tham khảo", level=1)
    references = [
        "[1] README.md, docs/00-project-brief.md, docs/02-current-system.md và docs/requirements/register.md, workspace AGVproject, đối chiếu ngày 20/10/2026.",
        "[2] docs/03-web.md, docs/04-server.md, docs/05-algorithms.md, docs/07-network.md, docs/09-testing.md và docs/18-hardening.md.",
        "[3] contracts/web-server.md và các ví dụ hợp đồng MQTT, HTTP, WebSocket trong thư mục contracts.",
        "[4] src/core/AGVserver.cpp, các header trong include, AGVwebserver/gateway.js, AGVwebserver/db.js và AGVwebserver/public/dashboard.js.",
        "[5] tests/README.md, hồ sơ kiểm thử phần mềm và các giới hạn bằng chứng, đối chiếu ngày 20/10/2026.",
        "[6] OASIS. MQTT Version 3.1.1, 2014. https://docs.oasis-open.org/mqtt/mqtt/v3.1.1/os/mqtt-v3.1.1-os.html.",
        "[7] IETF. RFC 6455 The WebSocket Protocol, 2011. https://datatracker.ietf.org/doc/html/rfc6455.",
        "[8] Dijkstra, E. W. A note on two problems in connexion with graphs. Numerische Mathematik, 1959.",
    ]
    for item in references:
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.left_indent = Cm(0.8)
        paragraph.paragraph_format.first_line_indent = Cm(-0.8)
        paragraph.paragraph_format.space_after = Pt(3)
        set_run_font(paragraph.add_run(item), 11)
    heading_sizes = {"Title": 19, "Subtitle": 13, "Heading 1": 15, "Heading 2": 13}
    for paragraph in doc.paragraphs:
        if paragraph.style.name in heading_sizes:
            for run in paragraph.runs:
                set_run_font(
                    run,
                    heading_sizes[paragraph.style.name],
                    bold=paragraph.style.name in {"Title", "Heading 1", "Heading 2"},
                    italic=paragraph.style.name == "Subtitle",
                )

    doc.core_properties.title = "Báo cáo kiến trúc và điều phối hệ thống AGV trong kho"
    doc.core_properties.subject = "Kiến trúc Core, gateway, SQLite, dashboard, an toàn, kiểm thử và hướng phát triển"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    architecture()
    core_modules()
    gateway_modules()
    warehouse_map()
    dashboard_wireframe()
    sequence_flow()
    safety_flow()
    build_doc()

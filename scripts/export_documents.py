"""Export the two AGV Markdown documents to editable Word documents."""
from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUTPUTS = (
    ("docs/00-project-brief.md", "AGV_Workspace_Overview.docx"),
    ("docs/requirements/template.md", "AGV_Requirements_Template.docx"),
)


def set_font(style, size: float) -> None:
    style.font.name = "Arial"
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    fonts = style.element.get_or_add_rPr()
    rfonts = fonts.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        fonts.append(rfonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn("w:" + attr), "Arial")


def add_link(paragraph, label: str, url: str) -> None:
    relation = paragraph.part.relate_to(
        url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True
    )
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), relation)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "18466B")
    props.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.append(underline)
    run.append(props)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    link.append(run)
    paragraph._p.append(link)


def write_inline(paragraph, text: str) -> None:
    tokens = re.split(r"(\[[^\]]+\]\([^)]+\)|\*\*[^*]+\*\*|`[^`]+`)", text)
    for token in tokens:
        link = re.fullmatch(r"\[([^\]]+)\]\(([^)]+)\)", token)
        if link and link.group(2).startswith(("http://", "https://")):
            add_link(paragraph, link.group(1), link.group(2))
        elif link:
            paragraph.add_run(link.group(1))
        elif token.startswith("**") and token.endswith("**"):
            paragraph.add_run(token[2:-2]).bold = True
        elif token.startswith("`") and token.endswith("`"):
            paragraph.add_run(token[1:-1])
        else:
            paragraph.add_run(token)


def add_table(doc, lines: list[str]) -> None:
    rows = [[part.strip() for part in line.strip().strip("|").split("|")] for line in lines]
    rows = [row for row in rows if not all(re.fullmatch(r":?-{3,}:?", cell) for cell in row)]
    cols = len(rows[0])
    table = doc.add_table(rows=0, cols=cols)
    table.autofit = False
    widths = [4.4, 12.4] if cols == 2 else [3.2, 5.7, 7.9] if cols == 3 else [16.8 / cols] * cols
    if rows[0][0] == "Topic MQTT":
        widths = [4.7, 3.0, 9.1]
    for column, width in zip(table.columns, widths):
        column.width = Cm(width)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement("w:" + edge)
        for attr, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
            element.set(qn("w:" + attr), value)
        borders.append(element)
    table._tbl.tblPr.append(borders)
    for index, values in enumerate(rows):
        row = table.add_row()
        no_split = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(no_split)
        if index == 0:
            repeat = OxmlElement("w:tblHeader")
            row._tr.get_or_add_trPr().append(repeat)
        for cell, width, value in zip(row.cells, widths, values):
            cell.width = Cm(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props = cell._tc.get_or_add_tcPr()
            shade = OxmlElement("w:shd")
            shade.set(qn("w:fill"), "DCE8F0" if index == 0 else "F4F6F8" if index % 2 == 0 else "FFFFFF")
            props.append(shade)
            margins = OxmlElement("w:tcMar")
            for edge in ("top", "bottom", "left", "right"):
                item = OxmlElement("w:" + edge)
                item.set(qn("w:w"), "105")
                item.set(qn("w:type"), "dxa")
                margins.append(item)
            props.append(margins)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(2)
            paragraph.paragraph_format.space_before = Pt(2)
            paragraph.paragraph_format.line_spacing = 1.06
            write_inline(paragraph, value)
            for run in paragraph.runs:
                run.font.size = Pt(9)
                run.bold = index == 0
                run.font.color.rgb = RGBColor(0, 0, 0)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def build(source: Path, destination: Path) -> None:
    doc = Document()
    doc.core_properties.title = source.read_text(encoding="utf-8").splitlines()[0].lstrip("# ")
    doc.core_properties.author = "Nhóm dự án AGV"
    doc.core_properties.subject = "Workspace nền nghiên cứu robot tự hành trong kho"
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(1.85), Cm(1.85)
    section.left_margin, section.right_margin = Cm(2.1), Cm(2.1)
    section.header_distance, section.footer_distance = Cm(0.8), Cm(0.8)
    for name, size in (("Normal", 11), ("Title", 21), ("Heading 1", 14), ("Heading 2", 12), ("List Bullet", 11), ("List Number", 11)):
        set_font(doc.styles[name], size)
    for border in list(doc.styles.element.iter(qn("w:pBdr"))):
        border.getparent().remove(border)
    for name in ("Title", "Heading 1", "Heading 2"):
        doc.styles[name].font.underline = False
    normal = doc.styles["Normal"].paragraph_format
    normal.space_after = Pt(7)
    normal.line_spacing = 1.15
    doc.styles["Title"].paragraph_format.space_after = Pt(14)
    for name in ("Heading 1", "Heading 2"):
        fmt = doc.styles[name].paragraph_format
        fmt.space_before, fmt.space_after = Pt(12), Pt(6)
        fmt.keep_with_next = True
    header = section.header.paragraphs[0]
    header.add_run("AGV IoT    Workspace nghiên cứu").font.size = Pt(8)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("Bộ nền 0.1    •    Trang ").font.size = Pt(8)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    lines = source.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if line.startswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i].strip())
                i += 1
            add_table(doc, table_lines)
            continue
        if line.startswith("# "):
            write_inline(doc.add_paragraph(style="Title"), line[2:])
        elif line.startswith("## "):
            paragraph = doc.add_paragraph(style="Heading 1")
            if source.name == "template.md" and line[3:] in (
                "Kho bản đồ và lộ trình", "Yêu cầu có thể nghiệm thu"
            ):
                paragraph.paragraph_format.page_break_before = True
            write_inline(paragraph, line[3:])
        elif line.startswith("### "):
            write_inline(doc.add_paragraph(style="Heading 2"), line[4:])
        elif line.startswith("- "):
            write_inline(doc.add_paragraph(style="List Bullet"), line[2:])
        elif re.match(r"^\d+\. ", line):
            write_inline(doc.add_paragraph(style="List Number"), re.sub(r"^\d+\. ", "", line))
        else:
            text = [line]
            while i + 1 < len(lines) and lines[i + 1].strip() and not re.match(r"^(#|\||- |\d+\. )", lines[i + 1].strip()):
                i += 1
                text.append(lines[i].strip())
            write_inline(doc.add_paragraph(), " ".join(text))
        i += 1
    for border in list(doc.element.iter(qn("w:pBdr"))):
        border.getparent().remove(border)
    destination.parent.mkdir(parents=True, exist_ok=True)
    doc.save(destination)
    print(destination.relative_to(ROOT))


def main() -> None:
    for source, filename in OUTPUTS:
        build(ROOT / source, ROOT / "docs" / "exports" / filename)


if __name__ == "__main__":
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    main()
